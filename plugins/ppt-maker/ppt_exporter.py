"""Editable PPTX export for ppt-maker."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


class PptxExportError(RuntimeError):
    """Raised when PPTX export fails."""


class PptxExporter:
    """Render slide IR into editable PowerPoint shapes."""

    def export(self, slides_ir: dict[str, Any], output_path: str | Path) -> Path:
        path = Path(output_path)
        path.parent.mkdir(parents=True, exist_ok=True)
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        theme = self._theme(slides_ir)
        for slide_ir in slides_ir.get("slides", []):
            self._render_slide(prs, slide_ir, theme)
        if len(prs.slides) == 0:
            raise PptxExportError("Cannot export an empty deck")
        prs.save(path)
        return path

    def _render_slide(self, prs: Presentation, slide_ir: dict[str, Any], theme: dict[str, str]) -> None:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        self._title(slide, slide_ir.get("title", "Untitled"), theme)
        slide_type = slide_ir.get("slide_type", "content")
        content = slide_ir.get("content", {})
        if slide_type in {"cover", "closing"}:
            self._center_text(slide, content.get("body") or slide_ir.get("notes", ""), theme)
        elif slide_type == "metric_cards":
            self._metric_cards(slide, content.get("metrics", []), theme)
            self._bullets(slide, content.get("bullets", []), top=4.6, theme=theme)
        elif slide_type in {"chart_bar", "chart_line", "chart_pie"}:
            self._simple_chart(slide, content.get("chart_spec") or {}, theme)
            self._bullets(slide, content.get("bullets", []), top=5.3, theme=theme)
        elif slide_type == "data_table":
            self._table(slide, content.get("rows", []), content.get("columns", []), theme)
        elif slide_type == "insight_summary":
            self._bullets(slide, content.get("findings", []), theme=theme)
            self._bullets(slide, content.get("risks", []), left=7.2, top=1.55, theme=theme, title="风险提示")
        else:
            bullets = content.get("bullets") or [content.get("body", "")]
            self._bullets(slide, bullets, theme=theme)

    def _title(self, slide, text: str, theme: dict[str, str]) -> None:
        box = slide.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(12.1), Inches(0.75))
        paragraph = box.text_frame.paragraphs[0]
        paragraph.text = text
        paragraph.font.size = Pt(28)
        paragraph.font.bold = True
        paragraph.font.name = theme["font_heading"]
        paragraph.font.color.rgb = self._rgb(theme["primary_color"])

    def _center_text(self, slide, text: str, theme: dict[str, str]) -> None:
        box = slide.shapes.add_textbox(Inches(1.2), Inches(2.7), Inches(10.9), Inches(1.4))
        paragraph = box.text_frame.paragraphs[0]
        paragraph.text = text
        paragraph.alignment = PP_ALIGN.CENTER
        paragraph.font.size = Pt(24)
        paragraph.font.name = theme["font_body"]

    def _bullets(
        self,
        slide,
        bullets: list[str],
        *,
        left: float = 0.9,
        top: float = 1.45,
        theme: dict[str, str],
        title: str | None = None,
    ) -> None:
        if title:
            title_box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(5.2), Inches(0.4))
            title_box.text_frame.text = title
            title_box.text_frame.paragraphs[0].font.bold = True
            top += 0.5
        box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(5.7), Inches(4.8))
        frame = box.text_frame
        frame.clear()
        for index, bullet in enumerate([item for item in bullets if item]):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.text = str(bullet)
            paragraph.level = 0
            paragraph.font.size = Pt(18)
            paragraph.font.name = theme["font_body"]

    def _metric_cards(self, slide, metrics: list[str], theme: dict[str, str]) -> None:
        values = metrics or ["Metric 1", "Metric 2", "Metric 3"]
        for index, metric in enumerate(values[:4]):
            left = 0.9 + index * 3.05
            shape = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Inches(left),
                Inches(1.8),
                Inches(2.65),
                Inches(1.5),
            )
            shape.fill.solid()
            shape.fill.fore_color.rgb = self._rgb(theme["primary_color"])
            shape.line.color.rgb = self._rgb(theme["primary_color"])
            frame = shape.text_frame
            frame.text = str(metric)
            frame.paragraphs[0].font.size = Pt(18)
            frame.paragraphs[0].font.bold = True
            frame.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)

    def _simple_chart(self, slide, chart_spec: dict[str, Any], theme: dict[str, str]) -> None:
        title = chart_spec.get("title") or "Chart suggestion"
        chart_box = slide.shapes.add_textbox(Inches(0.9), Inches(1.45), Inches(11.5), Inches(0.45))
        chart_box.text_frame.text = title
        bars = [0.65, 0.9, 0.55, 0.78]
        for index, value in enumerate(bars):
            left = 1.2 + index * 1.4
            height = 3.0 * value
            shape = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Inches(left),
                Inches(4.6 - height),
                Inches(0.8),
                Inches(height),
            )
            shape.fill.solid()
            shape.fill.fore_color.rgb = self._rgb(theme["accent_color"])
            shape.line.color.rgb = self._rgb(theme["accent_color"])

    def _table(self, slide, rows: list[list[str]], columns: list[str], theme: dict[str, str]) -> None:
        columns = columns[:6] or ["Column"]
        rows = rows[:8] or [[""] * len(columns)]
        table_shape = slide.shapes.add_table(
            len(rows) + 1,
            len(columns),
            Inches(0.7),
            Inches(1.4),
            Inches(11.8),
            Inches(4.8),
        )
        table = table_shape.table
        for col_index, column in enumerate(columns):
            table.cell(0, col_index).text = str(column)
        for row_index, row in enumerate(rows, start=1):
            for col_index, value in enumerate(row[: len(columns)]):
                table.cell(row_index, col_index).text = str(value)

    def _theme(self, slides_ir: dict[str, Any]) -> dict[str, str]:
        first = next(iter(slides_ir.get("slides", [])), {})
        theme = first.get("theme") or {}
        return {
            "primary_color": theme.get("primary_color") or "#3457D5",
            "secondary_color": theme.get("secondary_color") or "#172033",
            "accent_color": theme.get("accent_color") or "#FFB000",
            "font_heading": theme.get("font_heading") or "Microsoft YaHei",
            "font_body": theme.get("font_body") or "Microsoft YaHei",
        }

    def _rgb(self, color: str) -> RGBColor:
        cleaned = color.lstrip("#")
        if len(cleaned) != 6:
            cleaned = "3457D5"
        return RGBColor(int(cleaned[0:2], 16), int(cleaned[2:4], 16), int(cleaned[4:6], 16))

